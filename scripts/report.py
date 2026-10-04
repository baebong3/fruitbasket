"""리포트 생성 : 최신 스냅샷 + 누적 이력 -> HTML 대시보드, 품목 페이지, 마크다운, 엑셀

산출물
  docs/index.html                 최신 리포트 (GitHub Pages 첫 화면)
  docs/reports/YYYY-MM-DD.html    날짜별 보관본
  docs/items/<품목>-<품종>.html    품목별 상세 페이지
  docs/fruitbasket_prices.xlsx    엑셀 (최신 소매·도매, 최근 이력)
  reports/YYYY-MM-DD.md           GitHub에서 바로 읽는 요약
  reports/LATEST.md               최신 요약 사본
"""
from __future__ import annotations

import re
import sys
from collections import defaultdict
from datetime import date, timedelta, timezone
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import store  # noqa: E402
import items as itemmod  # noqa: E402
import kitchen  # noqa: E402
import basket as basketmod  # noqa: E402
from pricing import PriceBook  # noqa: E402
from theme import (GRADES, UP, DOWN, ACC, badge, chg_span, esc, face_svg, kdate, mdate, label_of,  # noqa: E402
                   pct, shell, sp, split_name, won, no_dash)

KST = timezone(timedelta(hours=9))
DOCS = store.ROOT / "docs"
REPORTS = store.ROOT / "reports"


# ---------- 계산 ----------

def grade_of(p: Decimal | None, g: dict) -> str | None:
    return itemmod.grade_for(p, g)


def fill_from_history(p: dict, sid: str, asof: str, hist_idx: dict, used: dict) -> dict:
    """KAMIS가 비워 둔 1주·1개월 전 가격(명절 휴장 등)을 저장된 이력의 가장 가까운 이전 조사일 값으로 채움"""
    series = hist_idx.get(sid)
    if not series:
        return p
    p = dict(p)
    base = date.fromisoformat(asof)
    for key, days in (("w1", 7), ("m1", 30)):
        if p.get(key) is not None:
            continue
        for back in range(0, 7):  # 목표일부터 최대 6일 이전까지
            d = (base - timedelta(days=days + back)).isoformat()
            if d in series:
                p[key] = series[d]
                used[key].add(d)
                break
    return p


def build_records(snap: dict, cfg: dict, history: list[dict] | None = None) -> tuple[dict, dict]:
    """({cls_name: {cat_name: [record,...]}}, 대체 비교일 {w1: set, m1: set})"""
    prio = cfg.get("rank_priority") or ["상품"]
    hist_idx: dict = defaultdict(dict)
    for h in history or []:
        hist_idx[store.series_id(h)][h["date"]] = h["price"]
    used: dict = {"w1": set(), "m1": set()}
    out: dict = defaultdict(dict)
    for g in snap["groups"]:
        # 품목·품종별 대표 등급 하나만
        groups: dict = defaultdict(list)
        for it in g["items"]:
            if it["prices"].get("today") is not None:
                groups[(it["item_code"], it["kind_code"])].append(it)
        recs = []
        for (ic, kc), cands in groups.items():
            it = itemmod.pick_rank(cands, prio)
            sid = "|".join([g["cls_code"], ic, kc, it["rank_code"]])
            p = fill_from_history(it["prices"], sid, snap["date"], hist_idx, used)
            name, kind = split_name(it["item_name"], it["kind_name"])
            base_key = "avg" if p.get("avg") else "y1"
            r = {
                "label": label_of(name, kind), "name": name, "kind": kind,
                "rank": it.get("rank", ""), "unit": it["unit"],
                "cat": g["cat_name"], "cls": g["cls_name"], "sid": sid,
                "item_code": ic, "kind_code": kc,
                "today": p.get("today"),
                "d1": pct(p.get("today"), p.get("d1")), "w1": pct(p.get("today"), p.get("w1")),
                "m1": pct(p.get("today"), p.get("m1")), "y1": pct(p.get("today"), p.get("y1")),
                "avg": pct(p.get("today"), p.get("avg")), "base_key": base_key,
            }
            r["base"] = r[base_key]
            r["grade"] = grade_of(r["base"], cfg["grade"])
            recs.append(r)
        if recs:
            out[g["cls_name"]][g["cat_name"]] = recs
    return out, used


# ---------- 요약 ----------

def summarize(recs_by_cls: dict) -> dict:
    retail = recs_by_cls.get("소매") or next(iter(recs_by_cls.values()), {})
    allr = [r for rs in retail.values() for r in rs]
    fruit = retail.get("과일류", allr)
    with_base = [r for r in fruit if r["base"] is not None]
    pricey = [r for r in with_base if r["grade"] in ("ORANGE", "RED")]
    cheap = sorted([r for r in with_base if r["grade"] == "GREEN"], key=lambda r: r["base"])
    top = max(with_base, key=lambda r: r["base"], default=None)
    low = min(with_base, key=lambda r: r["base"], default=None)
    cls_name = "소매" if "소매" in recs_by_cls else next(iter(recs_by_cls), "")
    cat_name = "과일" if "과일류" in retail else "품목"
    if top and top["base"] > 0:
        ref = "평년" if top["base_key"] == "avg" else "1년 전"
        head = (f"{top['label']} {ref}보다 {top['base']:,.1f}% 비싸, "
                f"{cls_name} {cat_name} {len(with_base):,}개 중 {len(pricey):,}개가 평년보다 비쌈")
    elif low:
        head = f"{cls_name} {cat_name} 가격 대체로 평년 이하, {low['label']} 평년보다 {abs(low['base']):,.1f}% 쌈"
    else:
        head = "가격 비교 기준(평년·1년 전) 자료 없음"
    up = sum(1 for r in allr if r["d1"] is not None and r["d1"] > 0)
    down = sum(1 for r in allr if r["d1"] is not None and r["d1"] < 0)
    return {"headline": head, "n_items": len(allr), "up": up, "down": down, "pricey": len(pricey),
            "n_fruit": len(with_base), "top": top, "low": low, "cheap": cheap[:6], "cls_name": cls_name}


# ---------- HTML 조각 ----------

def item_href(r: dict, rel: str, links: dict) -> str | None:
    p = links.get((r["item_code"], r["kind_code"]))
    return f"{rel}{p}" if p else None


def bars_html(recs: list[dict], rel: str, links: dict) -> str:
    rs = sorted([r for r in recs if r["base"] is not None], key=lambda r: r["base"], reverse=True)
    if not rs:
        return '<p class="empty">비교 기준 가격이 없는 품목만 있음</p>'
    mx = max(abs(r["base"]) for r in rs) or Decimal(1)
    rows = []
    for r in rs:
        ratio = float(abs(r["base"]) / mx)
        color = GRADES[r["grade"]][1]
        val = f'<span class="v">{sp(r["base"])}</span>'
        bar = f'<i style="width:calc((100% - 4.6em) * {ratio:.4f});background:{color}"></i>'
        neg = f"{val}{bar}" if r["base"] < 0 else ""
        pos = f"{bar}{val}" if r["base"] >= 0 else ""
        mark = "" if r["base_key"] == "avg" else '<sup title="평년가 없음, 1년 전 대비">*</sup>'
        sub = f'<small>{esc(r["kind"])}</small>' if r["kind"] else ""
        href = item_href(r, rel, links)
        full = f'<span class="fl">{esc(r["label"])}{mark}</span><span class="mn">{esc(r["name"])}{mark}{sub}</span>'
        lab = f'<a href="{href}">{full}</a>' if href else full
        rows.append(f'<div class="bl">{lab}</div><div class="bt"><div class="neg">{neg}</div><div class="pos">{pos}</div></div>')
    note = '<p class="note">* 평년가가 없어 1년 전 가격과 비교</p>' if any(r["base_key"] != "avg" for r in rs) else ""
    return f'<div class="bars">{"".join(rows)}</div>{note}'


def table_html(recs: list[dict], rel: str, links: dict) -> str:
    rs = sorted(recs, key=lambda r: (r["base"] is None, -(r["base"] or 0)))
    head = ("<tr><th class='l'>품목</th><th class='u'>단위</th><th>오늘 가격(원)</th><th>평년 대비</th><th>가격 수준</th>"
            "<th>전일 대비</th><th>1주 전 대비</th><th>1개월 전 대비</th><th>1년 전 대비</th></tr>")
    body = []
    for r in rs:
        sub = " · ".join(x for x in (r["kind"], r["unit"]) if x)
        inner = f'<span class="fl">{esc(r["label"])}</span><span class="mn">{esc(r["name"])}</span>'
        href = item_href(r, rel, links)
        cell = f'<a href="{href}">{inner}</a>' if href else inner
        body.append(
            f'<tr><td class="l">{cell}<small class="lu">{esc(sub)}</small></td><td class="u">{esc(r["unit"])}</td>'
            f'<td><span class="n"><b>{won(r["today"])}</b></span></td><td>{chg_span(r["avg"])}</td><td>{badge(r["grade"], 16)}</td>'
            f'<td>{chg_span(r["d1"])}</td><td>{chg_span(r["w1"])}</td><td>{chg_span(r["m1"])}</td><td>{chg_span(r["y1"])}</td></tr>')
    return f'<div class="tw"><table><thead>{head}</thead><tbody>{"".join(body)}</tbody></table></div>'


def spark_svg(points: list[tuple[str, int]], color: str) -> str:
    W, H, L, R, T, B = 320, 140, 8, 8, 26, 40
    vals = [v for _, v in points]
    lo, hi = min(vals), max(vals)
    if hi == lo:
        hi, lo = hi + 1, lo - 1
    n = len(points)
    xs = [L + (W - L - R) * i / max(n - 1, 1) for i in range(n)]
    ys = [T + (H - T - B) * (1 - (v - lo) / (hi - lo)) for v in vals]
    path = " ".join(f"{'M' if i == 0 else 'L'}{x:.1f},{y:.1f}" for i, (x, y) in enumerate(zip(xs, ys)))
    d0, d1 = points[0][0][5:].replace("-", "."), points[-1][0][5:].replace("-", ".")

    def ly(y: float, neighbor: float) -> float:
        below = neighbor < y
        if below and y + 18 > H - B + 8:
            below = False
        if not below and y - 10 < 12:
            below = True
        return y + 18 if below else y - 10
    y0 = ly(ys[0], ys[1] if n > 1 else ys[0])
    y1 = ly(ys[-1], ys[-2] if n > 1 else ys[-1])
    area = f"{path} L{xs[-1]:.1f},{H - B + 6:.1f} L{xs[0]:.1f},{H - B + 6:.1f} Z"
    return (f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="가격 추이">'
            f'<path d="{area}" fill="{color}" opacity=".1"/>'
            f'<path d="{path}" fill="none" stroke="{color}" stroke-width="2.4" stroke-linejoin="round" stroke-linecap="round"/>'
            f'<circle cx="{xs[0]:.1f}" cy="{ys[0]:.1f}" r="3.4" fill="#fff" stroke="{color}" stroke-width="2"/>'
            f'<circle cx="{xs[-1]:.1f}" cy="{ys[-1]:.1f}" r="4.2" fill="{color}" stroke="#fff" stroke-width="2"/>'
            f'<text x="{L}" y="{y0:.1f}" class="sv">{vals[0]:,}</text>'
            f'<text x="{W - R}" y="{y1:.1f}" class="sv" text-anchor="end">{vals[-1]:,}</text>'
            f'<text x="{L}" y="{H - 4}" class="sd">{d0}</text><text x="{W - R}" y="{H - 4}" class="sd" text-anchor="end">{d1}</text></svg>')


def trends_html(recs_by_cat: dict, history: list[dict], cfg: dict, asof: str, rel: str, links: dict) -> str:
    start = (date.fromisoformat(asof) - timedelta(days=int(cfg.get("trend_days", 90)))).isoformat()
    series = defaultdict(list)
    for h in history:
        if start <= h["date"] <= asof:
            series[store.series_id(h)].append((h["date"], h["price"]))
    recs = [r for rs in recs_by_cat.values() for r in rs if r["base"] is not None]
    recs.sort(key=lambda r: abs(r["base"]), reverse=True)
    cards = []
    for r in recs:
        pts = sorted(series.get(r["sid"], []))
        if len(pts) < 2:
            continue
        g = GRADES[r["grade"]]
        href = item_href(r, rel, links) or "#"
        cards.append(f'<a class="tc" href="{href}"><div class="tt">{face_svg(r["grade"], 20)}<b>{esc(r["label"])}</b>'
                     f'<span class="tu">{esc(r["unit"])}</span><span class="tp" style="background:{g[2]};color:{g[3]}">평년 {sp(r["base"])}</span></div>'
                     f'{spark_svg(pts, g[1])}</a>')
        if len(cards) >= int(cfg.get("trend_items", 8)):
            break
    if not cards:
        return '<p class="empty">이력이 2일 이상 쌓이면 추이 그래프가 표시됨</p>'
    return f'<div class="tg">{"".join(cards)}</div>'


def page_html(snap: dict, recs_by_cls: dict, s: dict, history: list[dict], cfg: dict,
              archive: list[str], rel: str, links: dict, used: dict, recipes: list[dict], bk: dict) -> str:
    asof = snap["date"]
    g = cfg["grade"]
    rng = {"GREEN": f'{g["green_max"]:+.0f}% 이하', "YELLOW": f'{g["green_max"]:+.0f}~{g["yellow_max"]:+.0f}%',
           "ORANGE": f'{g["yellow_max"]:+.0f}% 초과', "RED": f'{g["orange_max"]:+.0f}% 초과'}
    legend = '<div class="legend">' + "".join(f'<span>{face_svg(k, 15)}{GRADES[k][0]} ({rng[k]})</span>' for k in GRADES) + "</div>"
    fb = []
    for key, nm in (("w1", "1주 전"), ("m1", "1개월 전")):
        if used[key]:
            fb.append(f"{nm} 가격이 없는 품목(휴장 등)은 {', '.join(mdate(d) for d in sorted(used[key]))} 가격과 비교")
    fb_note = f'<p class="note">{esc(" · ".join(fb))}</p>' if fb else ""

    inputs, labels, panels = [], [], []
    for i, (cls, cats) in enumerate(recs_by_cls.items()):
        inputs.append(f'<input type="radio" name="t" id="t{i}"{" checked" if i == 0 else ""}>')
        labels.append(f'<label for="t{i}">{esc(cls)} 가격</label>')
        sec = []
        for cat, recs in cats.items():
            sec.append(f'<div class="card"><h2>{esc(cat)} {esc(cls)}<small>{len(recs):,}개 품목 · 품목명을 누르면 상세 페이지</small></h2>'
                       f"<h3>평년 대비 등락률</h3>{legend}{bars_html(recs, rel, links)}"
                       f"<h3>품목별 가격</h3>{table_html(recs, rel, links)}{fb_note}</div>")
        sec.append(f'<div class="card"><h2>{esc(cls)} 가격 추이<small>최근 {int(cfg.get("trend_days", 90)):,}일, 평년 대비 변동이 큰 품목</small></h2>'
                   + trends_html(cats, history, cfg, asof, rel, links) + "</div>")
        panels.append(f'<section class="panel" id="p{i}">{"".join(sec)}</section>')

    top = s["top"]
    top_txt = f'{esc(top["label"])}<small>{sp(top["base"])}</small>' if top and top["base"] is not None else "-"
    chips = "".join(f'<a class="chip" href="{item_href(r, rel, links) or "#"}">{face_svg("GREEN", 16)}{esc(r["label"])} {sp(r["base"])}</a>'
                    for r in s["cheap"])
    chips_html = f'<div class="chips">지금 사기 좋은 과일 {chips}</div>' if chips else ""
    arch = "".join(f'<li><a href="{rel}reports/{d}.html">{kdate(d)}</a></li>' for d in archive[:60])
    body = f"""<section class="hero"><div><div class="eyebrow">Today's Market</div>
<h1>{esc(s['headline'])}</h1>{chips_html}</div>
<div class="kpis">
<div class="kpi"><div class="k">조사 품목({esc(s['cls_name'])})</div><div class="x">{s['n_items']:,}<small>개</small></div></div>
<div class="kpi"><div class="k">어제보다 오름 / 내림</div><div class="x"><span style="color:{UP}">{s['up']:,}</span><small>/</small><span style="color:{DOWN}">{s['down']:,}</span></div></div>
<div class="kpi"><div class="k">평년보다 비싼 과일</div><div class="x">{s['pricey']:,}<small>/ {s['n_fruit']:,}개</small></div></div>
<div class="kpi"><div class="k">가장 비싸진 품목</div><div class="x" style="font-size:19px;white-space:normal;word-break:keep-all">{top_txt}</div></div>
</div></section>
{basketmod.section_html(bk, rel, links, False)}
<div class="card" id="kitchen"><h2>오늘의 알뜰 요리<small>평년보다 싼 재료로 · <a href="{rel}recipes.html">레시피 {len(recipes):,}개 모두 보기 ›</a></small></h2>
<p class="lead">주재료가 평년보다 싼 요리 {sum(1 for r in recipes if r["pick"]):,}가지 중 재료비 절약률이 큰 순. 비용은 KAMIS 소매 가격으로 계산한 농산물 재료비이고, 계란·두부 등은 참고가</p>
<div class="rgrid">{"".join(kitchen.recipe_card(r, rel, links, False) for r in recipes[:3])}</div></div>
<div class="tabs">{''.join(inputs)}<div class="seg">{''.join(labels)}</div>{''.join(panels)}</div>
<div class="card"><h2>지난 리포트<small>날짜를 누르면 그날 리포트</small></h2><ul class="arch">{arch}</ul></div>"""
    fetched = snap.get("fetched_at", "")[:16].replace("T", " ")
    return shell(f"과일바구니 가격 리포트 {asof}", body, rel, asof, fetched)


# ---------- 마크다운·엑셀 ----------

def markdown(snap: dict, recs_by_cls: dict, s: dict) -> str:
    L = [f"# 과일바구니 가격 리포트 {snap['date']}", "", f"**{s['headline']}**", "",
         f"- 조사 품목({s['cls_name']}) {s['n_items']:,}개, 전일 대비 상승 {s['up']:,}개 / 하락 {s['down']:,}개",
         f"- 평년보다 비싼 과일 {s['pricey']:,}개 / {s['n_fruit']:,}개", ""]
    for cls, cats in recs_by_cls.items():
        for cat, recs in cats.items():
            L += [f"## {cat} {cls}", "", "| 품목 | 단위 | 오늘(원) | 전일 | 1주 전 | 1개월 전 | 평년 | 수준 |", "|---|---|--:|--:|--:|--:|--:|:-:|"]
            for r in sorted(recs, key=lambda r: (r["base"] is None, -(r["base"] or 0))):
                gname = GRADES[r["grade"]][0] if r["grade"] else "-"
                L.append(f"| {r['label']} | {r['unit']} | {won(r['today'])} | {sp(r['d1'])} | {sp(r['w1'])} | {sp(r['m1'])} | {sp(r['avg'])} | {gname} |")
            L.append("")
    L.append("자료: 한국농수산식품유통공사(aT) KAMIS 농산물유통정보 · 자동 생성")
    return no_dash("\n".join(L) + "\n")


def excel(path: Path, recs_by_cls: dict, history: list[dict], asof: str, cfg: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    bold, thick, thin = Font(bold=True), Side(style="medium"), Side(style="thin", color="D9D9D9")
    for cls, cats in recs_by_cls.items():
        ws = wb.create_sheet(f"최신_{cls}")
        ws.append(["부류", "품목", "등급", "단위", "오늘 가격(원)", "전일 대비(%)", "1주 전 대비(%)",
                   "1개월 전 대비(%)", "1년 전 대비(%)", "평년 대비(%)", "가격 수준"])
        for c in ws[1]:
            c.font, c.border, c.alignment = bold, Border(bottom=thick), Alignment(horizontal="center")
        for cat, recs in cats.items():
            for r in recs:
                ws.append([cat, r["label"], r["rank"], r["unit"], r["today"]]
                          + [float(r[k]) if r[k] is not None else None for k in ("d1", "w1", "m1", "y1", "avg")]
                          + [GRADES[r["grade"]][0] if r["grade"] else None])
        for row in ws.iter_rows(min_row=2):
            row[4].number_format = "#,##0"
            for c in row[5:10]:
                c.number_format = "+0.0;-0.0;0.0"
            for c in row:
                c.border = Border(bottom=thin)
        for i, w in enumerate([8, 24, 8, 10, 13, 12, 13, 14, 13, 12, 10], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "C2"
    start = (date.fromisoformat(asof) - timedelta(days=int(cfg.get("trend_days", 90)))).isoformat()
    hist = [h for h in history if start <= h["date"] <= asof]
    dates = sorted({h["date"] for h in hist})
    meta, grid = {}, defaultdict(dict)
    for h in hist:
        sid = store.series_id(h)
        meta[sid] = h
        grid[sid][h["date"]] = h["price"]
    ws = wb.create_sheet("최근이력")
    ws.append(["구분", "부류", "품목", "품종", "등급", "단위"] + dates)
    for c in ws[1]:
        c.font, c.border = bold, Border(bottom=thick)
    for sid, h in sorted(meta.items(), key=lambda kv: (kv[1]["cls_code"], kv[1]["cat_code"], kv[1]["item_code"])):
        ws.append([h["cls_name"], h["cat_name"], h["item_name"], h["kind_name"], h["rank"], h["unit"]] + [grid[sid].get(d) for d in dates])
    for row in ws.iter_rows(min_row=2, min_col=7):
        for c in row:
            c.number_format = "#,##0"
    ws.freeze_panes = "G2"
    ws2 = wb.create_sheet("자료출처")
    ws2.append(["한국농수산식품유통공사(aT) KAMIS 농산물유통정보 일별 부류별 소매·도매 가격"])
    ws2.append(["가격 수준: 평년(최근 5년 중 최대·최소 제외 평균) 대비 등락률 기준"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------- 검증 게이트 ----------

def verify(html_text: str, n_bars: int | None = None, name: str = "index") -> list[str]:
    errs = []
    if "—" in html_text or "–" in html_text:
        errs.append(f"{name}: 줄표 포함")
    vals = re.findall(r'<span class="v">([^<]*)</span>', html_text)
    if n_bars is not None and len(vals) != n_bars:
        errs.append(f"{name}: 막대 수치 라벨 수 불일치 {len(vals)} != {n_bars}")
    if any(not v.strip() for v in vals):
        errs.append(f"{name}: 빈 막대 수치 라벨")
    texts = (re.findall(r'<span class="n"[^>]*>(?:<b>)?([^<]*)', html_text)
             + re.findall(r'class="(?:sv|lab|lab sub|ax)"[^>]*>([^<]*)<', html_text)
             + re.findall(r'<div class="big">([^<]*)<', html_text))
    for txt in texts:
        if re.fullmatch(r"\d{4}\.\d{2}", txt.strip()):  # 그래프 x축 연·월 표기
            continue
        if re.search(r"\d{4,}", re.sub(r"\.\d+", "", txt)):
            errs.append(f"{name}: 천 단위 콤마 누락 {txt!r}")
            break
    if re.search(r'class="lab[^"]*"[^>]*>\s*<', html_text):
        errs.append(f"{name}: 빈 그래프 라벨")
    return errs


def main() -> int:
    cfg = store.load_config()
    asof = sys.argv[1] if len(sys.argv) > 1 else store.latest_snapshot_date()
    if not asof:
        print("스냅샷 없음, 먼저 collect.py 실행")
        return 1
    snap = store.load_snapshot(asof)
    history = store.read_history()
    recs, used = build_records(snap, cfg, history)
    if not recs:
        print("표시할 품목 없음")
        return 1
    s = summarize(recs)
    fetched = snap.get("fetched_at", "")[:16].replace("T", " ")

    REPORTS.mkdir(parents=True, exist_ok=True)
    (DOCS / "reports").mkdir(parents=True, exist_ok=True)
    archive = sorted({p.stem for p in (DOCS / "reports").glob("*.html")} | {asof}, reverse=True)

    links, n_items = itemmod.build_item_pages(snap, cfg, history, fetched)
    book = PriceBook(snap, cfg, itemmod.series_index(history))
    recipes = kitchen.compute_recipes(book, cfg["grade"])
    bk = basketmod.build(book, history, asof)
    idx = page_html(snap, recs, s, history, cfg, archive, "", links, used, recipes, bk)
    arc = page_html(snap, recs, s, history, cfg, archive, "../", links, used, recipes, bk)
    picks = [r for r in recipes if r["pick"]]
    others = [r for r in recipes if not r["pick"]]
    rec_body = (f'<div class="card"><h2>오늘의 알뜰 요리<small>{esc(kdate(asof))} KAMIS 소매 가격 기준</small></h2>'
                f'<p class="lead">주재료가 평년보다 싼 요리 {len(picks):,}가지. 재료비는 조사 가격 × 레시피 수량으로 계산하고, '
                f'계란·두부·고기 등 조사하지 않는 재료는 참고가(data/staples.yml)로 더함. 기본 양념은 비용에서 제외</p>'
                f'<div class="rgrid full">{"".join(kitchen.recipe_card(r, "", links, True) for r in picks)}</div></div>'
                + (f'<div class="card"><h2>오늘은 평년 수준이거나 비싼 재료<small>주재료 가격 수준 보통 이상</small></h2>'
                   f'<div class="rgrid full">{"".join(kitchen.recipe_card(r, "", links, True) for r in others)}</div></div>' if others else ""))
    rec_html = shell("오늘의 알뜰 요리 - 과일바구니", rec_body, "", asof, fetched)
    bk_html = shell("과일바구니 장바구니 지수", basketmod.section_html(bk, "", links, True), "", asof, fetched)
    n_bars = sum(1 for cats in recs.values() for rs in cats.values() for r in rs if r["base"] is not None)
    errs = verify(idx, n_bars) + verify(rec_html, None, "recipes") + verify(bk_html, None, "basket")
    for p in sorted((DOCS / "items").glob("*.html")):
        errs += verify(p.read_text(encoding="utf-8"), None, p.stem)
    if errs:
        print("검증 실패:\n  " + "\n  ".join(errs[:20]))
        return 2

    (DOCS / "index.html").write_text(idx, encoding="utf-8")
    (DOCS / "reports" / f"{asof}.html").write_text(arc, encoding="utf-8")
    (DOCS / "recipes.html").write_text(rec_html, encoding="utf-8")
    (DOCS / "basket.html").write_text(bk_html, encoding="utf-8")
    (DOCS / ".nojekyll").write_text("", encoding="utf-8")
    md = markdown(snap, recs, s)
    (REPORTS / f"{asof}.md").write_text(md, encoding="utf-8")
    (REPORTS / "LATEST.md").write_text(md, encoding="utf-8")
    excel(DOCS / "fruitbasket_prices.xlsx", recs, history, asof, cfg)
    print(f"리포트 생성 {asof}: {s['headline']} (품목 페이지 {n_items:,}개, 알뜰 요리 {len(picks):,}개, 장바구니 지수 {bk['idx_avg_all']})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
